"""The final 35-university official-source list (Step 5A).

This file is the pilot's scope and its acquisition targets, so the tests are about the
two ways importing it could quietly go wrong: resolving a university to the wrong
target, and losing a claimed source responsibility to deduplication. Both failures are
silent, and both would be discovered months later by a consultant quoting one
university's fees under another's name.

The happy-path tests run against the **client's real file** and skip when it is not
present, because its actual shape — 385 URL cells, 319 distinct, 66 repeats, 3
genuinely new additional sources — is the thing worth asserting. The rejection tests
build deliberately broken copies of it, so what they prove is that the real importer
refuses real mistakes rather than that a synthetic fixture round-trips.
"""

from __future__ import annotations

import shutil
import uuid
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from sqlalchemy import Connection, text

from app.db.enums import UNCLASSIFIED_SOURCE_TYPE, SourceCandidateState, SourceCategory
from app.domains.onboarding.importer import import_target_list
from app.domains.pilot import queue, verification
from app.domains.pilot.official_sources import (
    ADDITIONAL_COLUMN,
    COLUMN_CATEGORIES,
    PILOT_INSTITUTION_COUNT,
    OfficialSourceListError,
    import_official_source_list,
)
from tests.integration.conftest import (
    OFFICIAL_SOURCES_ADDITIONAL_CELLS,
    OFFICIAL_SOURCES_CORE_CELLS,
    OFFICIAL_SOURCES_DISTINCT_PAGES,
    OFFICIAL_SOURCES_INSTITUTIONS,
    expect_violation,
)

pytestmark = pytest.mark.integration


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def targets(conn: Connection, client_workbook: Path) -> None:
    """The client's QS target list, imported. Nothing resolves without it."""
    import_target_list(conn, client_workbook)


@pytest.fixture
def imported(conn: Connection, targets: None, official_sources_workbook: Path) -> Any:
    return import_official_source_list(conn, official_sources_workbook)


def _edited(source: Path, tmp_path: Path, name: str) -> Any:
    """A writable copy of the client's file, plus a handle on its sheet."""
    copy = tmp_path / name
    shutil.copy(source, copy)
    book = openpyxl.load_workbook(copy)
    sheet = book["Universities"]
    header = {cell.value: cell.column for cell in sheet[1] if isinstance(cell.value, str)}
    return book, sheet, header, copy


# ===========================================================================
# 1-4. Scope and matching
# ===========================================================================


def test_the_final_list_resolves_exactly_thirty_five_institutions(
    conn: Connection, imported: Any
) -> None:
    """Requirement 1 and 3, against the client's own file.

    This is the regression test for the pilot's size: if the file, the constant and
    the database ever disagree, this fails rather than a later report quietly
    reporting 34.
    """
    report = imported
    assert PILOT_INSTITUTION_COUNT == 35
    assert report.institutions_in_file == OFFICIAL_SOURCES_INSTITUTIONS
    assert report.institutions_resolved == OFFICIAL_SOURCES_INSTITUTIONS
    assert report.institutions_ambiguous == 0
    assert report.institutions_unknown == 0

    staged = conn.execute(
        text("SELECT count(*) FROM pilot_selected_university WHERE is_selected")
    ).scalar_one()
    assert staged == OFFICIAL_SOURCES_INSTITUTIONS
    waved = conn.execute(
        text("SELECT count(*) FROM target_institution WHERE pilot_wave = 1")
    ).scalar_one()
    assert waved == OFFICIAL_SOURCES_INSTITUTIONS


def test_the_pilot_is_not_inferred_beyond_the_file(conn: Connection, imported: Any) -> None:
    """The other 146 targets keep `pilot_wave` NULL and stay valid future targets."""
    rows = conn.execute(
        text("SELECT pilot_wave, count(*) FROM target_institution GROUP BY 1 ORDER BY 1")
    ).all()
    counts = {row[0]: row[1] for row in rows}
    assert counts[1] == OFFICIAL_SOURCES_INSTITUTIONS
    assert counts[None] == 181 - OFFICIAL_SOURCES_INSTITUTIONS
    assert set(counts) == {1, None}, "a wave other than 1 was invented"


def test_a_near_miss_name_is_not_fuzzily_matched(
    conn: Connection, targets: None, official_sources_workbook: Path, tmp_path: Path
) -> None:
    """The failure this design refuses, using a name a similarity match would accept.

    `Imperial College` is a prefix of `Imperial College London` and shares every
    token. Any threshold loose enough to join them also joins `University of
    Canterbury` to `Canterbury Christ Church University`, and that mistake publishes
    one university's fees under another's name — silently, and permanently.
    """
    book, sheet, header, copy = _edited(official_sources_workbook, tmp_path, "near.xlsx")
    sheet.cell(row=2, column=header["University Name"], value="Imperial College")
    book.save(copy)

    with pytest.raises(OfficialSourceListError) as excinfo:
        import_official_source_list(conn, copy)
    problems = " ".join(excinfo.value.problems)
    assert "Imperial College" in problems
    assert "does not match" in problems
    assert conn.execute(text("SELECT count(*) FROM pilot_submission")).scalar_one() == 0


def test_an_unknown_university_stops_the_whole_import(
    conn: Connection, targets: None, official_sources_workbook: Path, tmp_path: Path
) -> None:
    """All or nothing. Importing the 34 that matched leaves a pilot that looks
    complete and is quietly missing a university."""
    book, sheet, header, copy = _edited(official_sources_workbook, tmp_path, "unknown.xlsx")
    sheet.cell(row=2, column=header["University Name"], value="University of Nowhere")
    book.save(copy)

    with pytest.raises(OfficialSourceListError):
        import_official_source_list(conn, copy)
    for table in ("pilot_submission", "pilot_selected_university", "pilot_collected_source"):
        assert conn.execute(text(f"SELECT count(*) FROM {table}")).scalar_one() == 0
    assert (
        conn.execute(
            text("SELECT count(*) FROM target_institution WHERE pilot_wave IS NOT NULL")
        ).scalar_one()
        == 0
    )


def test_a_repeated_university_row_is_rejected(
    conn: Connection, targets: None, official_sources_workbook: Path, tmp_path: Path
) -> None:
    """Two rows for one institution would give it two source sets and no way to say
    which is current."""
    book, sheet, header, copy = _edited(official_sources_workbook, tmp_path, "dupe.xlsx")
    name = sheet.cell(row=2, column=header["University Name"]).value
    sheet.cell(row=3, column=header["University Name"], value=name)
    book.save(copy)

    with pytest.raises(OfficialSourceListError) as excinfo:
        import_official_source_list(conn, copy)
    assert any("already appears on row" in problem for problem in excinfo.value.problems)


def test_a_country_disagreement_is_reported_not_resolved(
    conn: Connection, targets: None, official_sources_workbook: Path, tmp_path: Path
) -> None:
    """If the file says Australia and the target list says GB, one is wrong and the
    importer must not pick."""
    book, sheet, header, copy = _edited(official_sources_workbook, tmp_path, "region.xlsx")
    sheet.cell(row=2, column=header["Country/Region"], value="Australia")
    book.save(copy)

    with pytest.raises(OfficialSourceListError) as excinfo:
        import_official_source_list(conn, copy)
    assert any("one of them is wrong" in p.lower() for p in excinfo.value.problems)


def test_a_file_of_the_wrong_size_is_refused(
    conn: Connection, targets: None, official_sources_workbook: Path, tmp_path: Path
) -> None:
    """34 rows is a question for the client, never a row to invent."""
    book, sheet, header, copy = _edited(official_sources_workbook, tmp_path, "short.xlsx")
    sheet.delete_rows(2)
    book.save(copy)

    with pytest.raises(OfficialSourceListError) as excinfo:
        import_official_source_list(conn, copy)
    assert any("fixed at 35" in problem for problem in excinfo.value.problems)


# ===========================================================================
# 5-6. Idempotence and lineage
# ===========================================================================


def test_the_same_file_twice_writes_nothing(
    conn: Connection, imported: Any, official_sources_workbook: Path
) -> None:
    again = import_official_source_list(conn, official_sources_workbook)
    assert again.already_imported
    assert again.submission_id == imported.submission_id
    assert conn.execute(text("SELECT count(*) FROM pilot_submission")).scalar_one() == 1


def test_a_changed_file_becomes_a_new_submission(
    conn: Connection, imported: Any, official_sources_workbook: Path, tmp_path: Path
) -> None:
    """And the previous one is superseded, not overwritten."""
    book, sheet, header, copy = _edited(official_sources_workbook, tmp_path, "v2.xlsx")
    sheet.cell(row=2, column=header["Notes"], value="Revised note.")
    book.save(copy)

    second = import_official_source_list(conn, copy)
    assert second.submission_id is not None
    assert imported.submission_id is not None
    assert second.submission_id != imported.submission_id
    assert imported.submission_id in second.superseded_submission_ids

    statuses: dict[uuid.UUID, str] = {
        row.id: row.import_status
        for row in conn.execute(text("SELECT id, import_status FROM pilot_submission"))
    }
    assert statuses[imported.submission_id] == "SUPERSEDED"
    assert statuses[second.submission_id] == "VALIDATED"
    # The first submission's rows are untouched, which is what makes the two
    # comparable at all.
    assert (
        conn.execute(
            text("SELECT count(*) FROM pilot_collected_source WHERE submission_id = :s"),
            {"s": imported.submission_id},
        ).scalar_one()
        == imported.responsibilities
    )


def test_the_submission_records_that_it_defines_scope(conn: Connection, imported: Any) -> None:
    row = conn.execute(
        text(
            "SELECT submission_kind, defines_pilot_scope, expected_university_count, notes "
            "FROM pilot_submission WHERE id = :s"
        ),
        {"s": imported.submission_id},
    ).one()
    assert row.submission_kind == "OFFICIAL_SOURCE_LIST"
    assert row.defines_pilot_scope is True
    assert row.expected_university_count == PILOT_INSTITUTION_COUNT
    assert "Final pilot source list" in row.notes


def test_a_collection_workbook_cannot_claim_pilot_scope(conn: Connection) -> None:
    """Otherwise a facts workbook that merely omitted a university would drop it."""
    with expect_violation(conn, "only_a_source_list_defines_scope"):
        conn.execute(
            text(
                "INSERT INTO pilot_submission (id, file_sha256, original_filename, "
                "file_byte_size, template_version, submission_kind, defines_pilot_scope, "
                "selected_university_count) VALUES (:id, :sha, 'facts.xlsx', 1, 'v1', "
                "'COLLECTION_WORKBOOK', true, 0)"
            ),
            {"id": uuid.uuid4(), "sha": "a" * 64},
        )


# ===========================================================================
# 7-10. The workbook creates no evidence and no canonical fact
# ===========================================================================


def test_every_candidate_starts_pending(conn: Connection, imported: Any) -> None:
    states = conn.execute(
        text("SELECT DISTINCT verification_state FROM pilot_collected_source")
    ).scalars()
    assert set(states) == {SourceCandidateState.PENDING.value}


def test_the_import_creates_no_evidence_and_no_canonical_row(
    conn: Connection, imported: Any
) -> None:
    """Requirements 8, 9 and 10 in one sweep.

    An official-source list is a list of pages to look at later. Until one is fetched
    and snapshotted there is nothing to claim, and this asserts the importer took no
    shortcut toward pretending otherwise.
    """
    for table in (
        "field_claim",
        "field_provenance",
        "extraction",
        "snapshot",
        "content_blob",
        "fetch_run",
        "source",
        "source_mapping",
        "university",
        "program",
        "program_offering",
        "tuition",
        "admission_requirement",
        "language_requirement",
        "application_deadline",
        "change_proposal",
    ):
        count = conn.execute(text(f"SELECT count(*) FROM {table}")).scalar_one()
        assert count == 0, f"the source list wrote {count} row(s) to {table}"


def test_no_official_name_is_claimed_from_the_list_label(conn: Connection, imported: Any) -> None:
    """The workbook's University Name is the client's list label, not a collected
    official name. `university.name_en` is governed by C27, so putting a QS string
    into a column meaning "what the institution calls itself" would be the first step
    of exactly the laundering C27 exists to stop."""
    names = conn.execute(
        text("SELECT DISTINCT official_name_en FROM pilot_selected_university")
    ).scalars()
    assert set(names) == {None}
    homepages = conn.execute(
        text("SELECT count(*) FROM pilot_selected_university WHERE official_homepage IS NOT NULL")
    ).scalar_one()
    assert homepages == OFFICIAL_SOURCES_INSTITUTIONS


# ===========================================================================
# 11-15. Column mapping, the additional source, and deduplication
# ===========================================================================


def test_each_standard_column_maps_to_its_source_category(conn: Connection, imported: Any) -> None:
    """Requirement 11. Every core column produces one claim per institution."""
    counts = {
        row.source_type: row.n
        for row in conn.execute(
            text(
                "SELECT source_type, count(*) AS n FROM pilot_collected_source "
                "GROUP BY source_type"
            )
        )
    }
    for heading, category in COLUMN_CATEGORIES:
        assert (
            counts.get(category) == OFFICIAL_SOURCES_INSTITUTIONS
        ), f"{heading} -> {category} produced {counts.get(category)} rows"
        by_column = conn.execute(
            text(
                "SELECT count(*) FROM pilot_collected_source "
                "WHERE workbook_column = :c AND source_type = :t"
            ),
            {"c": heading, "t": category},
        ).scalar_one()
        assert by_column == OFFICIAL_SOURCES_INSTITUTIONS

    assert imported.core_url_cells == OFFICIAL_SOURCES_CORE_CELLS
    assert imported.additional_url_cells == OFFICIAL_SOURCES_ADDITIONAL_CELLS


def test_the_additional_source_is_optional(
    conn: Connection, targets: None, official_sources_workbook: Path, tmp_path: Path
) -> None:
    """Requirement 12. A file with the column blank still imports."""
    book, sheet, header, copy = _edited(official_sources_workbook, tmp_path, "noextra.xlsx")
    column = header[ADDITIONAL_COLUMN]
    for row in range(2, sheet.max_row + 1):
        # Attribute assignment, not cell(..., value=None): openpyxl reads the latter
        # as "no value argument" and leaves the cell untouched.
        sheet.cell(row=row, column=column).value = None
    book.save(copy)

    report = import_official_source_list(conn, copy)
    assert report.institutions_resolved == OFFICIAL_SOURCES_INSTITUTIONS
    assert report.additional_url_cells == 0
    assert report.unclassified_sources == 0
    assert report.responsibilities == OFFICIAL_SOURCES_CORE_CELLS


def test_a_repeated_additional_url_creates_no_second_page(conn: Connection, imported: Any) -> None:
    """Requirement 13. 32 of the 35 additional cells repeat a core URL.

    The row is kept — the client did supply it, and the lineage is worth having — but
    it is marked as a repeat, so it is not a second thing to fetch.
    """
    repeats = conn.execute(
        text(
            "SELECT count(*) FROM pilot_collected_source "
            "WHERE workbook_column = :c AND duplicate_of_source_ref IS NOT NULL"
        ),
        {"c": ADDITIONAL_COLUMN},
    ).scalar_one()
    assert repeats > 0

    # Every repeat points at a row holding the same URL for the same institution.
    mismatched = conn.execute(
        text(
            """
            SELECT count(*)
              FROM pilot_collected_source child
              JOIN pilot_collected_source parent
                ON parent.submission_id = child.submission_id
               AND parent.source_ref = child.duplicate_of_source_ref
             WHERE child.duplicate_of_source_ref IS NOT NULL
               AND (parent.url_sha256 <> child.url_sha256
                    OR parent.target_institution_id <> child.target_institution_id)
            """
        )
    ).scalar_one()
    assert mismatched == 0


def test_a_duplicated_url_keeps_every_claimed_responsibility(
    conn: Connection, imported: Any
) -> None:
    """Requirement 14, and the reason the old UNIQUE had to go.

    One admissions page answers for several categories. Deduplicating by URL would
    have discarded the claim that it is also the deadlines page, which is the fact
    the acquisition and review layers both need.
    """
    shared = conn.execute(
        text(
            """
            SELECT target_institution_id, url_sha256,
                   count(*) AS claims,
                   count(DISTINCT source_type) AS categories,
                   count(*) FILTER (WHERE duplicate_of_source_ref IS NULL) AS pages
              FROM pilot_collected_source
             GROUP BY 1, 2
            HAVING count(*) > 1
             ORDER BY claims DESC
            """
        )
    ).all()
    assert shared, "the real file repeats URLs; this test is vacuous without one"
    for row in shared:
        assert row.pages == 1, "a repeated URL produced more than one physical page"
        assert row.claims > 1, "the repeat was dropped instead of kept"

    assert (
        imported.responsibilities == OFFICIAL_SOURCES_CORE_CELLS + OFFICIAL_SOURCES_ADDITIONAL_CELLS
    )
    assert imported.physical_sources == OFFICIAL_SOURCES_DISTINCT_PAGES
    assert imported.duplicate_responsibilities == (
        imported.responsibilities - OFFICIAL_SOURCES_DISTINCT_PAGES
    )

    stored_pages = conn.execute(
        text("SELECT count(*) FROM pilot_collected_source WHERE duplicate_of_source_ref IS NULL")
    ).scalar_one()
    assert stored_pages == OFFICIAL_SOURCES_DISTINCT_PAGES
    assert (
        conn.execute(
            text("SELECT count(DISTINCT url_sha256) FROM pilot_collected_source")
        ).scalar_one()
        == OFFICIAL_SOURCES_DISTINCT_PAGES
    )


def test_two_physical_rows_for_one_url_are_refused(conn: Connection, imported: Any) -> None:
    """The partial unique index, tested directly: a second *unmarked* row for the
    same URL would become a second acquisition target for the same page."""
    row = conn.execute(
        text(
            "SELECT submission_id, target_institution_id, official_url, normalized_url, "
            "url_sha256, host FROM pilot_collected_source "
            "WHERE duplicate_of_source_ref IS NULL LIMIT 1"
        )
    ).one()
    with expect_violation(conn, "ix_pilot_collected_source_physical|duplicate key"):
        conn.execute(
            text(
                "INSERT INTO pilot_collected_source (submission_id, source_ref, "
                "target_institution_id, sheet_row_no, source_type, official_url, "
                "normalized_url, url_sha256, host) "
                "VALUES (:s, 'S9999', :t, 2, 'TUITION_FEES', :u, :n, :h, :host)"
            ),
            {
                "s": row.submission_id,
                "t": row.target_institution_id,
                "u": row.official_url,
                "n": row.normalized_url,
                "h": row.url_sha256,
                "host": row.host,
            },
        )


def test_a_distinct_additional_source_stays_unclassified(conn: Connection, imported: Any) -> None:
    """Requirement 15. Three of the additional URLs are genuinely new.

    Nothing in the workbook says what they are, and the column they arrived in is not
    evidence: one of them is a fee-rates table, not the PDF the heading hints at. They
    wait for a person.
    """
    rows = conn.execute(
        text(
            "SELECT source_ref, official_url FROM pilot_collected_source "
            "WHERE source_type = :t AND duplicate_of_source_ref IS NULL"
        ),
        {"t": UNCLASSIFIED_SOURCE_TYPE},
    ).all()
    assert len(rows) == imported.unclassified_distinct_sources
    assert 0 < len(rows) < OFFICIAL_SOURCES_INSTITUTIONS

    # And nothing was invented for them, including the tempting OFFICIAL_PDF.
    assert (
        conn.execute(
            text(
                "SELECT count(*) FROM pilot_collected_source " "WHERE source_type = 'OFFICIAL_PDF'"
            )
        ).scalar_one()
        == 0
    )


def test_an_unclassified_page_cannot_be_verified(conn: Connection, imported: Any) -> None:
    """Verification asserts a page is authoritative *for something*."""
    candidate = conn.execute(
        text("SELECT id FROM pilot_collected_source WHERE source_type = :t LIMIT 1"),
        {"t": UNCLASSIFIED_SOURCE_TYPE},
    ).scalar_one()
    actor = verification.Actor(user_id=_actor(conn))

    with pytest.raises(verification.CandidateDecisionRefusedError, match="no category yet"):
        verification.verify_candidate(
            conn, candidate_id=candidate, actor=actor, reason="looks official to me"
        )
    # The database refuses it too, so bypassing the service changes nothing.
    with expect_violation(conn, "unclassified_is_not_verifiable"):
        conn.execute(
            text(
                "UPDATE pilot_collected_source SET verification_state = 'VERIFIED', "
                "verified_at = now(), verified_by = :a, verification_reason = 'x' "
                "WHERE id = :id"
            ),
            {"id": candidate, "a": actor.user_id},
        )


def test_an_unclassified_page_may_be_rejected_and_then_classified(
    conn: Connection, imported: Any
) -> None:
    """Non-vacuity for the test above, and the route out of UNCLASSIFIED.

    Rejecting needs no category — "not a page we want" is a complete thought. Getting
    to VERIFIED needs one, and supplying it is its own recorded decision.
    """
    actor = verification.Actor(user_id=_actor(conn))
    unclassified = (
        conn.execute(
            text(
                "SELECT id FROM pilot_collected_source WHERE source_type = :t "
                "AND duplicate_of_source_ref IS NULL ORDER BY source_ref"
            ),
            {"t": UNCLASSIFIED_SOURCE_TYPE},
        )
        .scalars()
        .all()
    )
    assert len(unclassified) >= 2

    verification.reject_candidate(
        conn, candidate_id=unclassified[0], actor=actor, reason="Marketing flyer, not a fee table."
    )

    verification.classify_candidate(
        conn,
        candidate_id=unclassified[1],
        source_category=SourceCategory.TUITION_FEES.value,
        actor=actor,
        reason="Graduate fee-rates table for the current year.",
    )
    verification.verify_candidate(
        conn, candidate_id=unclassified[1], actor=actor, reason="Official fee schedule."
    )
    row = conn.execute(
        text("SELECT source_type, verification_state FROM pilot_collected_source WHERE id = :id"),
        {"id": unclassified[1]},
    ).one()
    assert row.source_type == SourceCategory.TUITION_FEES.value
    assert row.verification_state == SourceCandidateState.VERIFIED.value

    actions = (
        conn.execute(
            text("SELECT action FROM audit_log WHERE object_id = :id ORDER BY seq"),
            {"id": unclassified[1]},
        )
        .scalars()
        .all()
    )
    assert actions == ["PILOT_SOURCE_CLASSIFY", "PILOT_SOURCE_VERIFY"]


def test_classifying_as_unclassified_is_refused(conn: Connection, imported: Any) -> None:
    candidate = conn.execute(
        text("SELECT id FROM pilot_collected_source WHERE source_type = :t LIMIT 1"),
        {"t": UNCLASSIFIED_SOURCE_TYPE},
    ).scalar_one()
    with pytest.raises(verification.CandidateDecisionRefusedError, match="not a classification"):
        verification.classify_candidate(
            conn,
            candidate_id=candidate,
            source_category=UNCLASSIFIED_SOURCE_TYPE,
            actor=verification.Actor(user_id=_actor(conn)),
            reason="still unsure",
        )


# ===========================================================================
# 16-19. URLs, safety, and verification
# ===========================================================================


def test_raw_and_normalized_urls_are_both_kept(conn: Connection, imported: Any) -> None:
    """Requirement 16. Debugging a fetch that went to the wrong place needs both."""
    rows = conn.execute(
        text(
            "SELECT official_url, normalized_url, url_sha256, host "
            "FROM pilot_collected_source LIMIT 400"
        )
    ).all()
    assert rows
    for row in rows:
        assert row.official_url.startswith("https://")
        assert row.normalized_url.startswith("https://")
        assert len(row.url_sha256) == 64
        assert row.host == row.host.lower()
        assert row.host in row.normalized_url

    # Conservative normalisation: this file needed none, and nothing invented a
    # change. A rewritten path or a dropped query parameter addresses another page.
    unchanged = conn.execute(
        text("SELECT count(*) FROM pilot_collected_source WHERE official_url = normalized_url")
    ).scalar_one()
    assert unchanged == imported.responsibilities

    # The one query-addressed page keeps its parameters.
    query_urls = conn.execute(
        text("SELECT count(*) FROM pilot_collected_source WHERE normalized_url LIKE '%?%'")
    ).scalar_one()
    assert query_urls > 0


def test_a_non_http_url_is_rejected_and_not_stored(
    conn: Connection, targets: None, official_sources_workbook: Path, tmp_path: Path
) -> None:
    """Requirement 17. A workbook is not a way past URL safety (D23)."""
    book, sheet, header, copy = _edited(official_sources_workbook, tmp_path, "scheme.xlsx")
    sheet.cell(row=2, column=header["Tuition/Fees URL"], value="file:///etc/passwd")
    sheet.cell(
        row=3,
        column=header["Academic Calendar URL"],
        value="http://169.254.169.254/latest/meta-data/",
    )
    book.save(copy)

    report = import_official_source_list(conn, copy)
    assert len(report.rejected_urls) == 2
    assert (
        conn.execute(
            text("SELECT count(*) FROM pilot_collected_source WHERE official_url LIKE 'file:%'")
        ).scalar_one()
        == 0
    )
    assert (
        conn.execute(
            text("SELECT count(*) FROM pilot_collected_source WHERE host LIKE '169.254%'")
        ).scalar_one()
        == 0
    )
    # The other 383 still imported: one bad address does not cost the rest.
    assert (
        report.responsibilities
        == OFFICIAL_SOURCES_CORE_CELLS + OFFICIAL_SOURCES_ADDITIONAL_CELLS - 2
    )


def test_only_http_and_https_can_be_stored_at_all(conn: Connection, imported: Any) -> None:
    """The CHECK, not just the service: a fix-up script cannot introduce another
    scheme either."""
    row = conn.execute(
        text("SELECT submission_id, target_institution_id FROM pilot_collected_source LIMIT 1")
    ).one()
    with expect_violation(conn, "official_url_is_http"):
        conn.execute(
            text(
                "INSERT INTO pilot_collected_source (submission_id, source_ref, "
                "target_institution_id, sheet_row_no, source_type, official_url, "
                "normalized_url, url_sha256, host) VALUES (:s, 'S9998', :t, 2, "
                "'TUITION_FEES', 'ftp://example.ac.uk/fees', 'https://example.ac.uk/fees', "
                ":h, 'example.ac.uk')"
            ),
            {"s": row.submission_id, "t": row.target_institution_id, "h": "b" * 64},
        )


def test_a_matching_verified_host_does_not_auto_verify(conn: Connection, imported: Any) -> None:
    """Requirement 18. The strongest automatic signal available, acted on by nobody.

    Every one of these URLs is on a host the collector believes official, and 35 of
    them are the homepage itself. Verifying a real domain here still leaves every
    candidate PENDING.
    """
    institution, host = conn.execute(
        text(
            "SELECT target_institution_id, host FROM pilot_collected_source "
            "WHERE source_type = 'UNIVERSITY_HOME' ORDER BY source_ref LIMIT 1"
        )
    ).one()
    _verified_domain(conn, institution, host)

    queued = conn.execute(
        text(
            "SELECT host_matches_verified_domain, verification_state, is_physical_source "
            "FROM pilot_source_verification_queue "
            "WHERE target_institution_id = :t AND host = :h"
        ),
        {"t": institution, "h": host},
    ).all()
    assert queued
    assert all(row.host_matches_verified_domain for row in queued)
    assert {row.verification_state for row in queued} == {SourceCandidateState.PENDING.value}

    summary = queue.queue_summary(conn)
    assert summary.host_matches_verified_domain == len(queued)
    assert summary.overall.verified == 0


def test_verification_still_requires_actor_time_and_reason(conn: Connection, imported: Any) -> None:
    """Requirement 19, in both directions: the service refuses, and so does the CHECK."""
    candidate = conn.execute(
        text("SELECT id FROM pilot_collected_source WHERE source_type = 'TUITION_FEES' LIMIT 1")
    ).scalar_one()
    actor = verification.Actor(user_id=_actor(conn))

    with pytest.raises(verification.CandidateDecisionRefusedError, match="reason"):
        verification.verify_candidate(conn, candidate_id=candidate, actor=actor, reason="  ")

    with expect_violation(conn, "decision_records_who_when_why"):
        conn.execute(
            text(
                "UPDATE pilot_collected_source SET verification_state = 'VERIFIED' WHERE id = :id"
            ),
            {"id": candidate},
        )

    verification.verify_candidate(
        conn,
        candidate_id=candidate,
        actor=actor,
        reason="Institution-wide fee table, linked from the admissions page.",
    )
    row = conn.execute(
        text(
            "SELECT verified_at, verified_by, verification_reason "
            "FROM pilot_collected_source WHERE id = :id"
        ),
        {"id": candidate},
    ).one()
    assert row.verified_at is not None
    assert row.verified_by == actor.user_id
    assert "fee table" in row.verification_reason


# ===========================================================================
# 20. The queue report
# ===========================================================================


def test_the_queue_reports_identity_and_responsibility_separately(
    conn: Connection, imported: Any
) -> None:
    """Requirement 13-14's reporting half: a reviewer must see both.

    "385 claims" and "319 pages to fetch" are both true, and confusing them means
    either fetching a page three times or reviewing one claim out of three.
    """
    summary = queue.queue_summary(conn)
    assert summary.overall.total == imported.responsibilities
    assert summary.overall.pending == imported.responsibilities
    assert summary.physical_sources == OFFICIAL_SOURCES_DISTINCT_PAGES
    assert summary.distinct_urls == OFFICIAL_SOURCES_DISTINCT_PAGES
    assert summary.duplicate_responsibilities == (
        imported.responsibilities - OFFICIAL_SOURCES_DISTINCT_PAGES
    )
    assert summary.unclassified_sources == imported.unclassified_distinct_sources

    assert len(summary.by_institution) == OFFICIAL_SOURCES_INSTITUTIONS
    assert set(summary.by_source_type) == {category for _, category in COLUMN_CATEGORIES} | {
        UNCLASSIFIED_SOURCE_TYPE
    }
    for _, category in COLUMN_CATEGORIES:
        assert summary.by_source_type[category].total == OFFICIAL_SOURCES_INSTITUTIONS

    pages = queue.physical_sources(conn)
    assert len(pages) == OFFICIAL_SOURCES_DISTINCT_PAGES
    assert all(page.is_physical_source for page in pages)
    # Nothing is fetchable yet, which is the correct state and not an empty result.
    assert queue.physical_sources(conn, verified_only=True) == []


def test_the_report_counts_only_this_pilot_by_default(conn: Connection, imported: Any) -> None:
    """The other 146 target institutions have no candidates, so the two views agree
    here — but the filter is what stops a future non-pilot import inflating the
    worklist."""
    assert queue.open_count(conn) == imported.responsibilities
    assert queue.queue_summary(conn, selected_only=False).overall.total == (
        imported.responsibilities
    )


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------


def _actor(conn: Connection) -> uuid.UUID:
    existing: uuid.UUID | None = conn.execute(
        text("SELECT id FROM app_user WHERE email = 'sources-reviewer@example.test'")
    ).scalar_one_or_none()
    if existing is not None:
        return existing
    actor_id = uuid.uuid4()
    conn.execute(
        text(
            "INSERT INTO app_user (id, email, display_name) "
            "VALUES (:id, 'sources-reviewer@example.test', 'Sources Reviewer')"
        ),
        {"id": actor_id},
    )
    return actor_id


def _verified_domain(conn: Connection, target_institution_id: uuid.UUID, host: str) -> uuid.UUID:
    domain_id = uuid.uuid4()
    conn.execute(
        text(
            "INSERT INTO official_domain (id, target_institution_id, host, covers_subdomains, "
            "verification_status, verification_method, verification_evidence, verified_at, "
            "verified_by, is_active) VALUES (:id, :target, :host, true, 'VERIFIED_OFFICIAL', "
            "'GOVERNMENT_REGISTRY', 'Registry record', now(), :actor, true)"
        ),
        {
            "id": domain_id,
            "target": target_institution_id,
            "host": host,
            "actor": _actor(conn),
        },
    )
    return domain_id

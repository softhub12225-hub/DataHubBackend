"""Pilot collection readiness: selection, workbook-local references, scope, seeds.

The thing being defended here is that **the system never chooses the 36**. Everything
else follows from it: if selection is the client's decision, then the workbook must be
checkable rather than interpretable, which is why programmes and sources are joined by
explicit codes and why an unmappable applicant scope parks instead of defaulting.

Selection-count tests use synthetic institutions rather than the client's real
workbook, so they run everywhere. The one test that is genuinely about the client's
real distribution uses the real file and skips without it.
"""

from __future__ import annotations

import uuid
from pathlib import Path
from typing import Any

import pytest
from sqlalchemy import Connection, text

from app.domains.onboarding.importer import import_target_list
from app.domains.onboarding.pilot_template import (
    PROGRAM_REF_PATTERN,
    SOURCE_REF_PATTERN,
    build_pilot_template,
)
from app.domains.onboarding.pilot_workbook import (
    IssueCode,
    ScopeResolution,
    Severity,
    resolve_applicant_scope,
    validate_workbook,
)
from tests.integration.conftest import CLIENT_DISTRIBUTION

pytestmark = pytest.mark.integration


# ---------------------------------------------------------------------------
# Building a filled-in workbook
# ---------------------------------------------------------------------------


def make_candidates(conn: Connection, count: int, *, destination: str = "GB") -> list[uuid.UUID]:
    """`count` invented institutions in scope, ready to be exported."""
    list_id = uuid.uuid4()
    conn.execute(
        text(
            "INSERT INTO target_list (id, list_name, list_version, file_name, "
            "file_sha256, file_byte_size, sheet_name, imported_row_count) "
            "VALUES (:id, 'Synthetic List', :ver, 's.xlsx', :sha, 1, 's', :n)"
        ),
        {
            "id": list_id,
            "ver": list_id.hex[:8],
            "sha": f"{list_id.hex}{list_id.hex}",
            "n": count,
        },
    )
    ids: list[uuid.UUID] = []
    for index in range(count):
        target = uuid.uuid4()
        conn.execute(
            text(
                "INSERT INTO target_institution (id, match_key, first_seen_list_id, "
                "latest_list_id, destination_code) "
                "VALUES (:id, :key, :list, :list, :dest)"
            ),
            {
                "id": target,
                "key": f"synthetic-{index:04d}-{target.hex[:8]}",
                "list": list_id,
                "dest": destination,
            },
        )
        conn.execute(
            text(
                "INSERT INTO target_list_entry (target_list_id, target_institution_id, "
                "source_row, qs_name, qs_name_normalized, qs_rank, region_label, "
                "destination_code) VALUES (:list, :target, :row, :name, :norm, :rank, "
                "'英国', :dest)"
            ),
            {
                "list": list_id,
                "target": target,
                "row": index + 2,
                "name": f"Invented Institution {index:04d}",
                "norm": f"invented institution {index:04d}",
                "rank": index + 1,
                "dest": destination,
            },
        )
        ids.append(target)
    return ids


class Filler:
    """Writes values into an exported template, by column name."""

    def __init__(self, path: Path) -> None:
        import openpyxl

        self.path = path
        self.book = openpyxl.load_workbook(path)

    def _header(self, sheet: str) -> dict[str, int]:
        return {
            cell.value: cell.column for cell in self.book[sheet][1] if isinstance(cell.value, str)
        }

    def mark_selected(self, targets: list[uuid.UUID]) -> Filler:
        sheet = self.book["Pilot_Universities"]
        header = self._header("Pilot_Universities")
        wanted = {str(t) for t in targets}
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row=row, column=1).value in wanted:
                sheet.cell(row=row, column=header["selected"], value="YES")
                sheet.cell(row=row, column=header["official_name_en"], value="Invented University")
                sheet.cell(
                    row=row,
                    column=header["official_homepage"],
                    value="https://invented.example.ac.uk/",
                )
        return self

    def add(self, sheet_name: str, **values: Any) -> Filler:
        sheet = self.book[sheet_name]
        header = self._header(sheet_name)
        row = sheet.max_row + 1
        for name, value in values.items():
            if name not in header:
                raise KeyError(f"{sheet_name} has no column {name!r}")
            sheet.cell(row=row, column=header[name], value=value)
        return self

    def save(self) -> Path:
        self.book.save(self.path)
        return self.path


@pytest.fixture
def exported(conn: Connection, tmp_path: Path) -> Any:
    """40 candidates exported, so 35 / 36 / 37 are all expressible."""
    targets = make_candidates(conn, 40)
    out = tmp_path / "pilot.xlsx"
    build_pilot_template(conn, out, pilot_target=36)
    return targets, out


# ===========================================================================
# 1-2. Export: all candidates, none chosen
# ===========================================================================


def test_all_fifty_seven_pilot_candidates_are_exported(
    conn: Connection, tmp_path: Path, client_workbook: Path
) -> None:
    """Validation 1. The real GB/HK/MO distribution, from the client's own file."""
    import openpyxl

    import_target_list(conn, client_workbook)
    out = tmp_path / "real.xlsx"
    build_pilot_template(conn, out, destination_codes=["GB", "HK", "MO"], pilot_target=36)

    sheet = openpyxl.load_workbook(out)["Pilot_Universities"]
    expected = sum(CLIENT_DISTRIBUTION[code] for code in ("GB", "HK", "MO"))
    assert expected == 57
    assert sheet.max_row - 1 == 57


def test_the_export_chooses_none_automatically(conn: Connection, exported: Any) -> None:
    """Validation 2. Not by rank, not by row order, not by `--pilot-target`."""
    import openpyxl

    _targets, out = exported
    sheet = openpyxl.load_workbook(out)["Pilot_Universities"]
    header = {c.value: c.column for c in sheet[1]}
    selected = [
        sheet.cell(row=row, column=header["selected"]).value for row in range(2, sheet.max_row + 1)
    ]
    assert set(selected) == {None}, "the export pre-selected something"


def test_export_assigns_no_pilot_wave(conn: Connection, exported: Any) -> None:
    """`--pilot-target 36` is a validation figure and must write nothing."""
    _targets, _out = exported
    assert (
        conn.execute(
            text("SELECT count(*) FROM target_institution WHERE pilot_wave IS NOT NULL")
        ).scalar_one()
        == 0
    )


def test_pilot_target_does_not_reach_the_selection_column(conn: Connection, tmp_path: Path) -> None:
    """Non-vacuity: exporting with and without a target produces the same selection."""
    import openpyxl

    make_candidates(conn, 5)
    with_target = tmp_path / "a.xlsx"
    without = tmp_path / "b.xlsx"
    build_pilot_template(conn, with_target, pilot_target=3)
    build_pilot_template(conn, without)

    def selection(path: Path) -> list[Any]:
        sheet = openpyxl.load_workbook(path)["Pilot_Universities"]
        header = {c.value: c.column for c in sheet[1]}
        return [
            sheet.cell(row=row, column=header["selected"]).value
            for row in range(2, sheet.max_row + 1)
        ]

    assert selection(with_target) == selection(without) == [None] * 5


# ===========================================================================
# 3-5. Selection count
# ===========================================================================


@pytest.mark.parametrize(
    ("marked", "expected", "should_pass"),
    [(36, 36, True), (35, 36, False), (37, 36, False)],
)
def test_the_selected_count_is_validated_not_corrected(
    conn: Connection, exported: Any, marked: int, expected: int, should_pass: bool
) -> None:
    """Validations 3, 4 and 5.

    Both directions are errors. Thirty-five is not "nearly right" to be topped up,
    and thirty-seven is not trimmed: choosing either way would be the system making
    the client's decision.
    """
    targets, out = exported
    Filler(out).mark_selected(targets[:marked]).save()

    report = validate_workbook(conn, out, expected_selected=expected)
    assert len(report.selected_institutions) == marked

    mismatches = [
        issue for issue in report.issues if issue.code is IssueCode.SELECTED_COUNT_MISMATCH
    ]
    if should_pass:
        assert mismatches == []
        assert report.is_importable, [str(i) for i in report.errors]
    else:
        assert len(mismatches) == 1
        assert mismatches[0].severity is Severity.ERROR
        assert str(marked) in mismatches[0].message
        assert not report.is_importable


def test_a_selected_institution_must_carry_its_official_details(
    conn: Connection, exported: Any
) -> None:
    import openpyxl

    targets, out = exported
    Filler(out).mark_selected(targets[:36]).save()
    book = openpyxl.load_workbook(out)
    sheet = book["Pilot_Universities"]
    header = {c.value: c.column for c in sheet[1]}
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=header["selected"]).value == "YES":
            # Attribute assignment, NOT cell(..., value=None): openpyxl treats a
            # None value argument as "no value supplied" and leaves the cell as it
            # was, which would make this test pass vacuously.
            sheet.cell(row=row, column=header["official_homepage"]).value = None
            break
    book.save(out)

    report = validate_workbook(conn, out, expected_selected=36)
    assert any(i.code is IssueCode.SELECTED_WITHOUT_DETAILS for i in report.issues)


# ===========================================================================
# 6-9. Workbook-local references
# ===========================================================================


def _filled(conn: Connection, exported: Any, **overrides: Any) -> Any:
    """A minimal valid workbook: 36 selected, one programme, one source, one fee."""
    targets, out = exported
    first = str(targets[0])
    filler = Filler(out).mark_selected(targets[:36])
    filler.add(
        "Programs",
        target_institution_id=first,
        program_ref=overrides.get("program_ref", "P0001"),
        program_name_en="MSc Invented Computing",
        degree_level_code="MASTER",
        discipline_code="COMPUTER_AND_DATA",
        discipline_hint="Computer Science",
    )
    filler.add(
        "Official_Sources",
        target_institution_id=first,
        source_ref=overrides.get("source_ref", "S0001"),
        source_type="TUITION_FEES",
        official_url="https://invented.example.ac.uk/fees",
        checked_at="2027-01-15",
    )
    filler.add(
        "Tuition",
        target_institution_id=first,
        program_ref=overrides.get("tuition_program_ref", "P0001"),
        academic_year="2027/28",
        student_category_code="INTERNATIONAL",
        amount_kind="EXACT",
        amount_min=32000,
        amount_max=32000,
        currency_code="GBP",
        billing_unit_code="PER_YEAR",
        applicant_scope_hint=overrides.get("scope_hint", "UNIVERSAL"),
        official_text="GBP 32,000 per year",
        amount_status="PUBLISHED",
        source_ref=overrides.get("tuition_source_ref", "S0001"),
    )
    return filler, targets, out


def test_a_well_formed_workbook_validates(conn: Connection, exported: Any) -> None:
    """Non-vacuity guard for every rejection below."""
    filler, _targets, out = _filled(conn, exported)
    filler.save()
    report = validate_workbook(conn, out, expected_selected=36)
    assert report.is_importable, [str(i) for i in report.errors]
    assert report.program_refs.keys() == {"P0001"}
    assert report.source_refs.keys() == {"S0001"}


def test_a_duplicate_program_ref_is_rejected(conn: Connection, exported: Any) -> None:
    """Validation 6. The classic copied-row mistake."""
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Programs",
        target_institution_id=str(targets[0]),
        program_ref="P0001",
        program_name_en="MSc Invented Data Science",
        degree_level_code="MASTER",
        discipline_code="COMPUTER_AND_DATA",
        discipline_hint="Data Science",
    ).save()

    report = validate_workbook(conn, out, expected_selected=36)
    duplicates = [i for i in report.issues if i.code is IssueCode.PROGRAM_REF_DUPLICATE]
    assert len(duplicates) == 1
    assert "already used on row" in duplicates[0].message
    assert not report.is_importable


def test_an_unknown_program_ref_is_rejected(conn: Connection, exported: Any) -> None:
    """Validation 7. No fallback to matching the programme by name."""
    filler, _targets, out = _filled(conn, exported, tuition_program_ref="P0099")
    filler.save()

    report = validate_workbook(conn, out, expected_selected=36)
    unknown = [i for i in report.issues if i.code is IssueCode.PROGRAM_REF_UNKNOWN]
    assert len(unknown) == 1
    assert "P0099" in unknown[0].message
    assert unknown[0].sheet == "Tuition"
    assert not report.is_importable


def test_a_duplicate_source_ref_is_rejected(conn: Connection, exported: Any) -> None:
    """Validation 8."""
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Official_Sources",
        target_institution_id=str(targets[0]),
        source_ref="S0001",
        source_type="APPLICATION_DEADLINES",
        official_url="https://invented.example.ac.uk/deadlines",
        checked_at="2027-01-15",
    ).save()

    report = validate_workbook(conn, out, expected_selected=36)
    duplicates = [i for i in report.issues if i.code is IssueCode.SOURCE_REF_DUPLICATE]
    assert len(duplicates) == 1
    assert not report.is_importable


def test_an_unknown_source_ref_is_rejected(conn: Connection, exported: Any) -> None:
    """Validation 9. And emphatically not resolved by matching the URL."""
    filler, _targets, out = _filled(conn, exported, tuition_source_ref="S0099")
    filler.save()

    report = validate_workbook(conn, out, expected_selected=36)
    unknown = [i for i in report.issues if i.code is IssueCode.SOURCE_REF_UNKNOWN]
    assert len(unknown) == 1
    assert "S0099" in unknown[0].message
    assert not report.is_importable


def test_a_ref_belonging_to_another_institution_is_rejected(
    conn: Connection, exported: Any
) -> None:
    """The ambiguous cross-sheet case: the code exists, but not for this university."""
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Tuition",
        target_institution_id=str(targets[1]),
        program_ref="P0001",
        academic_year="2027/28",
        student_category_code="INTERNATIONAL",
        amount_kind="EXACT",
        amount_min=1000,
        amount_max=1000,
        currency_code="GBP",
        billing_unit_code="PER_YEAR",
        applicant_scope_hint="UNIVERSAL",
        official_text="GBP 1,000",
        amount_status="PUBLISHED",
        source_ref="S0001",
    ).save()

    report = validate_workbook(conn, out, expected_selected=36)
    codes = {i.code for i in report.issues}
    assert IssueCode.PROGRAM_REF_INSTITUTION_MISMATCH in codes
    assert IssueCode.SOURCE_REF_INSTITUTION_MISMATCH in codes


def test_a_malformed_ref_is_rejected(conn: Connection, exported: Any) -> None:
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Programs",
        target_institution_id=str(targets[0]),
        program_ref="see below",
        program_name_en="MSc Something",
        degree_level_code="MASTER",
        discipline_code="BUSINESS",
        discipline_hint="Finance",
    ).save()
    report = validate_workbook(conn, out, expected_selected=36)
    assert any(i.code is IssueCode.PROGRAM_REF_MALFORMED for i in report.issues)


def test_the_reference_patterns_are_what_the_readme_promises() -> None:
    import re

    assert re.match(PROGRAM_REF_PATTERN, "P0001")
    assert not re.match(PROGRAM_REF_PATTERN, "P1")
    assert re.match(SOURCE_REF_PATTERN, "S0042")
    assert not re.match(SOURCE_REF_PATTERN, "S0042 ")


# ===========================================================================
# 10. Applicant scope
# ===========================================================================


@pytest.mark.parametrize(
    "hint",
    [
        "China",
        "Mainland China",
        "Chinese bachelor degree",
        "985/211 institutions",
        "IB",
        "A-level",
        "international students",
        "",
        None,
        "  ",
    ],
)
def test_a_non_universal_scope_never_becomes_universal(hint: str | None) -> None:
    """Validation 10, at the unit level: only the literal UNIVERSAL resolves.

    A blank hint parks too. "The collector did not say" is not evidence that the
    university meant everybody.
    """
    assert resolve_applicant_scope(hint) is ScopeResolution.SCOPE_MAPPING_REQUIRED


@pytest.mark.parametrize("hint", ["UNIVERSAL", "universal", " Universal "])
def test_the_universal_hint_resolves(hint: str) -> None:
    """Non-vacuity guard: the one value that does resolve."""
    assert resolve_applicant_scope(hint) is ScopeResolution.UNIVERSAL


def test_a_scoped_requirement_is_parked_for_a_human(conn: Connection, exported: Any) -> None:
    """Validation 10, end to end through the workbook."""
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Language_Requirements",
        target_institution_id=str(targets[0]),
        program_ref="P0001",
        test_type_code="IELTS",
        overall_score=6.5,
        applicant_scope_hint="holders of a Chinese bachelor degree from a 985/211 institution",
        applicant_country_code="CN",
        qualification_hint="985/211",
        official_text="IELTS 7.0, or 6.5 for holders of a Chinese bachelor degree",
        requirement_status="PUBLISHED",
        source_ref="S0001",
    ).save()

    report = validate_workbook(conn, out, expected_selected=36)
    parked = [i for i in report.issues if i.code is IssueCode.SCOPE_MAPPING_REQUIRED]
    assert len(parked) == 1
    assert parked[0].severity is Severity.WARNING
    assert "NOT be treated as applying to everybody" in parked[0].message
    # A warning, so the workbook is still importable -- the row is held, not the file.
    assert report.is_importable, [str(i) for i in report.errors]
    assert report.scope_parked == 1


def test_the_scope_columns_are_present_on_every_fact_sheet(conn: Connection, exported: Any) -> None:
    import openpyxl

    _targets, out = exported
    book = openpyxl.load_workbook(out)
    for name in ("Admissions", "Language_Requirements", "Tuition", "Deadlines"):
        header = {c.value for c in book[name][1]}
        assert {
            "applicant_scope_hint",
            "applicant_country_code",
            "qualification_hint",
            "official_text",
        } <= header, f"{name} is missing scope columns"


# ===========================================================================
# 6 (field status) — blank vs OFFICIALLY_NOT_PUBLISHED
# ===========================================================================


def test_a_published_fact_without_a_source_is_rejected(conn: Connection, exported: Any) -> None:
    filler, _targets, out = _filled(conn, exported, tuition_source_ref=None)
    filler.save()
    report = validate_workbook(conn, out, expected_selected=36)
    assert any(i.code is IssueCode.SOURCE_REF_MISSING for i in report.issues)


def test_officially_not_published_needs_a_source_but_no_value(
    conn: Connection, exported: Any
) -> None:
    """The distinction the client asked to preserve, both halves of it.

    A checked absence still names the page it checked -- otherwise it is just a
    shrug -- but it must not require a fabricated amount.
    """
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Tuition",
        target_institution_id=str(targets[0]),
        program_ref="P0001",
        academic_year="2028/29",
        student_category_code="INTERNATIONAL",
        applicant_scope_hint="UNIVERSAL",
        official_text="Fees for 2028/29 have not yet been set.",
        amount_status="OFFICIALLY_NOT_PUBLISHED",
        source_ref="S0001",
    ).save()

    report = validate_workbook(conn, out, expected_selected=36)
    assert report.is_importable, [str(i) for i in report.errors]


def test_a_value_contradicting_its_status_is_rejected(conn: Connection, exported: Any) -> None:
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Tuition",
        target_institution_id=str(targets[0]),
        program_ref="P0001",
        academic_year="2029/30",
        student_category_code="INTERNATIONAL",
        amount_kind="EXACT",
        amount_min=1234,
        amount_max=1234,
        currency_code="GBP",
        billing_unit_code="PER_YEAR",
        applicant_scope_hint="UNIVERSAL",
        amount_status="OFFICIALLY_NOT_PUBLISHED",
        source_ref="S0001",
    ).save()

    report = validate_workbook(conn, out, expected_selected=36)
    assert any(i.code is IssueCode.VALUE_WITH_UNPUBLISHED_STATUS for i in report.issues)


def test_a_published_status_with_no_value_is_rejected(conn: Connection, exported: Any) -> None:
    """Never force a fabricated value -- but PUBLISHED must mean something is there."""
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Tuition",
        target_institution_id=str(targets[0]),
        program_ref="P0001",
        academic_year="2030/31",
        student_category_code="INTERNATIONAL",
        applicant_scope_hint="UNIVERSAL",
        amount_status="PUBLISHED",
        source_ref="S0001",
    ).save()

    report = validate_workbook(conn, out, expected_selected=36)
    published_without = [i for i in report.issues if i.code is IssueCode.PUBLISHED_WITHOUT_VALUE]
    assert len(published_without) == 1
    assert "do not invent a value" in published_without[0].message


def test_a_blank_row_asserts_nothing(conn: Connection, exported: Any) -> None:
    """Leaving a fact sheet empty is a valid answer and raises nothing."""
    targets, out = exported
    Filler(out).mark_selected(targets[:36]).save()
    report = validate_workbook(conn, out, expected_selected=36)
    assert report.is_importable, [str(i) for i in report.errors]
    assert report.rows_by_sheet.get("Tuition", 0) == 0


# ===========================================================================
# 11-12. Discipline seeds
# ===========================================================================


def test_the_three_pilot_disciplines_are_seeded(conn: Connection) -> None:
    """Validation 11."""
    rows = conn.execute(text("SELECT code, name_en, parent_id FROM discipline ORDER BY code")).all()
    assert [row.code for row in rows] == ["BUSINESS", "COMPUTER_AND_DATA", "ENGINEERING"]
    assert all(row.parent_id is None for row in rows), "the pilot seeds are top level only"


def test_no_deeper_taxonomy_was_invented(conn: Connection) -> None:
    """The client asked for three codes, not a subject tree."""
    assert conn.execute(text("SELECT count(*) FROM discipline")).scalar_one() == 3


def test_reseeding_disciplines_is_idempotent(conn: Connection) -> None:
    """Validation 12. `ON CONFLICT (code) DO NOTHING`, run twice."""
    before = conn.execute(text("SELECT count(*) FROM discipline")).scalar_one()
    for _ in range(2):
        conn.execute(
            text(
                "INSERT INTO discipline (code, name_en, name_zh) VALUES "
                "('BUSINESS', 'Business and Management', '商科与管理'), "
                "('COMPUTER_AND_DATA', 'Computer Science and Data', '计算机与数据'), "
                "('ENGINEERING', 'Engineering', '工程') "
                "ON CONFLICT (code) DO NOTHING"
            )
        )
    assert conn.execute(text("SELECT count(*) FROM discipline")).scalar_one() == before


def test_the_discipline_dropdown_offers_exactly_the_seeds(conn: Connection, exported: Any) -> None:
    import openpyxl

    _targets, out = exported
    book = openpyxl.load_workbook(out)
    lists = book["_Lists"]
    columns = {
        lists.cell(row=1, column=col).value: [
            lists.cell(row=row, column=col).value
            for row in range(2, lists.max_row + 1)
            if lists.cell(row=row, column=col).value is not None
        ]
        for col in range(1, lists.max_column + 1)
    }
    assert columns.get("discipline_code") == [
        "BUSINESS",
        "COMPUTER_AND_DATA",
        "ENGINEERING",
    ]


def test_the_discipline_hint_is_collected_verbatim(conn: Connection, exported: Any) -> None:
    """Source terminology is preserved for later mapping, not normalised away."""
    import openpyxl

    _targets, out = exported
    sheet = openpyxl.load_workbook(out)["Programs"]
    header = {c.value: c for c in sheet[1]}
    assert "discipline_hint" in header
    comment = header["discipline_hint"].comment
    assert comment is not None
    assert "exactly as you write it" in comment.text


# ===========================================================================
# README
# ===========================================================================


def test_the_readme_covers_every_required_topic(conn: Connection, exported: Any) -> None:
    """Requirement 7, asserted rather than assumed."""
    import openpyxl

    _targets, out = exported
    readme = "\n".join(
        str(row[0].value)
        for row in openpyxl.load_workbook(out)["README"].iter_rows()
        if row[0].value
    )
    for topic in (
        "WHICH COLUMNS ARE MANDATORY",
        "HOW TO MARK THE PILOT UNIVERSITIES",
        "PROGRAMME CODES (program_ref)",
        "SOURCE CODES (source_ref)",
        "BLANK IS NOT THE SAME AS 'NOT PUBLISHED'",
        "WHEN A RULE APPLIES ONLY TO SOME APPLICANTS",
        "Do not guess",
        "Copy official wording exactly",
        "Use only official sources",
    ):
        assert topic in readme, f"README does not cover: {topic}"


def test_no_collection_sheet_contains_example_data(conn: Connection, exported: Any) -> None:
    """Fabricated rows get returned as if they were collected. Examples live in the
    README only."""
    import openpyxl

    _targets, out = exported
    book = openpyxl.load_workbook(out)
    for name in (
        "Official_Sources",
        "Programs",
        "Admissions",
        "Language_Requirements",
        "Tuition",
        "Deadlines",
    ):
        assert book[name].max_row == 1, f"{name} ships with data rows"


def test_the_readme_examples_are_invented(conn: Connection, exported: Any) -> None:
    """The README may carry examples, but not real institutions."""
    import openpyxl

    readme = "\n".join(
        str(row[0].value)
        for row in openpyxl.load_workbook(exported[1])["README"].iter_rows()
        if row[0].value
    )
    for real in ("Imperial College", "University of Oxford", "Tsinghua", "HKU"):
        assert real not in readme, f"README names a real institution: {real}"


# ===========================================================================
# Problems found by adversarial review of the first draft
# ===========================================================================


@pytest.mark.parametrize("answer", ["Y", "是", "x", "TRUE", "1", "selected"])
def test_an_unrecognised_selection_answer_is_an_error_not_a_no(
    conn: Connection, exported: Any, answer: str
) -> None:
    """The first draft treated anything but YES as "not selected", silently.

    A collector writing Y or the Chinese for yes means yes. Dropping that institution
    from the pilot without saying so is the worst outcome for the one number the
    client controls, and the count would still have looked plausible.
    """
    import openpyxl

    targets, out = exported
    Filler(out).mark_selected(targets[:36]).save()
    book = openpyxl.load_workbook(out)
    sheet = book["Pilot_Universities"]
    header = {c.value: c.column for c in sheet[1]}
    # Row 38 is the 37th institution: not among the 36 already marked.
    sheet.cell(row=38, column=header["selected"], value=answer)
    book.save(out)

    report = validate_workbook(conn, out, expected_selected=36)
    unrecognised = [i for i in report.issues if i.code is IssueCode.SELECTED_VALUE_UNRECOGNISED]
    assert len(unrecognised) == 1, f"{answer!r} was silently ignored"
    assert unrecognised[0].severity is Severity.ERROR
    assert unrecognised[0].row == 38
    assert not report.is_importable


def test_a_language_requirement_with_only_subscores_is_publishable(
    conn: Connection, exported: Any
) -> None:
    """Requirement 6: never force a fabricated value.

    A page stating only "no component below 6.0" publishes a real requirement and has
    no overall figure. Demanding one would make the collector invent it.
    """
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Language_Requirements",
        target_institution_id=str(targets[0]),
        program_ref="P0001",
        test_type_code="IELTS",
        subscore_minimums="no component below 6.0",
        applicant_scope_hint="UNIVERSAL",
        official_text="No component below 6.0.",
        requirement_status="PUBLISHED",
        source_ref="S0001",
    ).save()

    report = validate_workbook(conn, out, expected_selected=36)
    assert not [i for i in report.issues if i.code is IssueCode.PUBLISHED_WITHOUT_VALUE], [
        str(i) for i in report.issues
    ]
    assert report.is_importable, [str(i) for i in report.errors]


def test_a_malformed_academic_year_is_caught_before_import(conn: Connection, exported: Any) -> None:
    """The academic-year CHECK is on two tables. Catch it while the client still has
    the page open, not halfway through a batch insert."""
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Tuition",
        target_institution_id=str(targets[0]),
        program_ref="P0001",
        academic_year="2027-28 entry",
        student_category_code="INTERNATIONAL",
        amount_kind="EXACT",
        amount_min=1,
        amount_max=1,
        currency_code="GBP",
        billing_unit_code="PER_YEAR",
        applicant_scope_hint="UNIVERSAL",
        official_text="x",
        amount_status="PUBLISHED",
        source_ref="S0001",
    ).save()

    report = validate_workbook(conn, out, expected_selected=36)
    bad = [i for i in report.issues if i.code is IssueCode.ACADEMIC_YEAR_MALFORMED]
    assert len(bad) == 1
    assert "2027-28 entry" in bad[0].message


def test_a_student_category_from_another_destination_is_caught(
    conn: Connection, tmp_path: Path
) -> None:
    """LOCAL is ambiguous between HK and MO by code alone, and nothing stopped a
    collector choosing a UK category for a Hong Kong institution."""
    targets = make_candidates(conn, 3, destination="HK")
    out = tmp_path / "hk.xlsx"
    build_pilot_template(conn, out)
    filler = Filler(out).mark_selected(targets)
    first = str(targets[0])
    filler.add(
        "Official_Sources",
        target_institution_id=first,
        source_ref="S0001",
        source_type="TUITION_FEES",
        official_url="https://invented.example.edu.hk/fees",
        checked_at="2027-01-15",
    )
    filler.add(
        "Programs",
        target_institution_id=first,
        program_ref="P0001",
        program_name_en="MSc Invented",
        degree_level_code="MASTER",
        discipline_code="BUSINESS",
        discipline_hint="Finance",
        study_mode="FULL_TIME",
        delivery_mode="ON_CAMPUS",
        duration_value=1,
        duration_unit="YEAR",
    )
    filler.add(
        "Tuition",
        target_institution_id=first,
        program_ref="P0001",
        academic_year="2027/28",
        student_category_code="HOME",
        amount_kind="EXACT",
        amount_min=100000,
        amount_max=100000,
        currency_code="HKD",
        billing_unit_code="PER_YEAR",
        applicant_scope_hint="UNIVERSAL",
        official_text="HKD 100,000",
        amount_status="PUBLISHED",
        source_ref="S0001",
    ).save()

    report = validate_workbook(conn, out, expected_selected=3)
    mismatched = [
        i for i in report.issues if i.code is IssueCode.STUDENT_CATEGORY_NOT_IN_DESTINATION
    ]
    assert len(mismatched) == 1
    assert "LOCAL and NON_LOCAL" in mismatched[0].message


@pytest.mark.parametrize(
    ("parts", "expected_fragment"),
    [
        ({"deadline_kind": "FIXED_DATE"}, "no year is given"),
        (
            {"deadline_kind": "FIXED_DATE", "year": 2027, "day": 15},
            "day was given without a month",
        ),
        (
            {"deadline_kind": "FIXED_DATE", "year": 2027, "month": 3, "time_of_day": "23:59"},
            "time was given without an exact day",
        ),
        (
            {
                "deadline_kind": "FIXED_DATE",
                "year": 2027,
                "month": 3,
                "day": 15,
                "timezone": "GMT",
            },
            "timezone was given without a time",
        ),
        (
            {
                "deadline_kind": "FIXED_DATE",
                "year": 2027,
                "month": 3,
                "day": 15,
                "month_part": "MID",
            },
            "cannot both be given",
        ),
        ({"deadline_kind": "NO_FIXED_DEADLINE", "year": 2027}, "there is no date to record"),
    ],
)
def test_inconsistent_deadline_parts_are_caught_before_import(
    conn: Connection, exported: Any, parts: dict[str, Any], expected_fragment: str
) -> None:
    """The D17 rules are database CHECKs the client cannot see.

    Without this the workbook validates clean and then throws constraint violations
    partway through the batch -- and each of these is a case where accepting the row
    would have meant manufacturing precision the page never gave.
    """
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Deadlines",
        target_institution_id=str(targets[0]),
        program_ref="P0001",
        academic_year="2027/28",
        intake_season_code="SEPTEMBER",
        deadline_text_verbatim="as published",
        applicant_scope_hint="UNIVERSAL",
        official_text="as published",
        deadline_status="PUBLISHED",
        source_ref="S0001",
        **parts,
    ).save()

    report = validate_workbook(conn, out, expected_selected=36)
    inconsistent = [i for i in report.issues if i.code is IssueCode.DEADLINE_PARTS_INCONSISTENT]
    assert inconsistent, [str(i) for i in report.issues]
    assert any(expected_fragment in i.message for i in inconsistent)


def test_a_consistent_deadline_passes(conn: Connection, exported: Any) -> None:
    """Non-vacuity guard for the six refusals above."""
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Deadlines",
        target_institution_id=str(targets[0]),
        program_ref="P0001",
        academic_year="2027/28",
        intake_season_code="SEPTEMBER",
        round_code="MAIN",
        deadline_kind="FIXED_DATE",
        deadline_text_verbatim="15 January 2027, 23:59 GMT",
        year=2027,
        month=1,
        day=15,
        time_of_day="23:59",
        timezone="GMT",
        applicant_scope_hint="UNIVERSAL",
        official_text="15 January 2027, 23:59 GMT",
        deadline_status="PUBLISHED",
        source_ref="S0001",
    ).save()

    report = validate_workbook(conn, out, expected_selected=36)
    assert not [i for i in report.issues if i.code is IssueCode.DEADLINE_PARTS_INCONSISTENT], [
        str(i) for i in report.issues
    ]


def test_an_unresolved_row_does_not_also_report_a_bogus_unknown_ref(
    conn: Connection, exported: Any
) -> None:
    """Without an institution there is no scope to resolve a ref against, and a
    second error on the same row sends the operator to the wrong cell."""
    filler, targets, out = _filled(conn, exported)
    filler.add(
        "Tuition",
        target_institution_id="not-a-uuid",
        program_ref="P0001",
        academic_year="2027/28",
        student_category_code="INTERNATIONAL",
        amount_kind="EXACT",
        amount_min=1,
        amount_max=1,
        currency_code="GBP",
        billing_unit_code="PER_YEAR",
        applicant_scope_hint="UNIVERSAL",
        official_text="x",
        amount_status="PUBLISHED",
        source_ref="S0001",
    ).save()

    report = validate_workbook(conn, out, expected_selected=36)
    tuition_issues = [i for i in report.issues if i.sheet == "Tuition"]
    codes = {i.code for i in tuition_issues}
    assert IssueCode.TARGET_UNRESOLVED in codes
    assert IssueCode.PROGRAM_REF_UNKNOWN not in codes
    assert IssueCode.SOURCE_REF_UNKNOWN not in codes


def test_the_offering_dimensions_are_collected_not_invented() -> None:
    """`program_offering` needs four NOT NULL columns that also form its identity.

    Collected as optional, an importer would invent them -- and the invented values
    would BE the identity, so the next workbook that filled them in would mint a
    second offering with a second full set of fees.
    """
    from app.domains.onboarding.pilot_template import PROGRAMS

    required = {c.name for c in PROGRAMS.columns if c.requirement == "REQUIRED"}
    assert {"study_mode", "delivery_mode", "duration_value", "duration_unit"} <= required


def test_the_source_degree_level_help_matches_what_absence_means(
    conn: Connection, exported: Any
) -> None:
    """The first draft said "leave blank if it serves all of them", inverting
    `SourceDegreeScope`: absence means NOT covered, so a diligently filled workbook
    would have produced a coverage report reading zero across the board."""
    import openpyxl

    _targets, out = exported
    sheet = openpyxl.load_workbook(out)["Official_Sources"]
    comment = next(c.comment.text for c in sheet[1] if c.value == "degree_level" and c.comment)
    assert "NOT that it covers everyone" in comment
    assert "one row per group" in comment
